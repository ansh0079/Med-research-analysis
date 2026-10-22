import json
import openpyxl

wb = openpyxl.load_workbook('Signal_MD_All_Topics_250_Per_Sheet.xlsx', data_only=True)

for sheet_name in ['Topics 0251-0500', 'Topics 1001-1250', 'Topics 1501-1750', 'Topics 1751-1875']:
    ws = wb[sheet_name]
    topics = []
    for r in range(7, ws.max_row + 1):
        no_val = ws.cell(row=r, column=1).value
        topic_val = ws.cell(row=r, column=2).value
        status_val = ws.cell(row=r, column=11).value
        if no_val:
            topics.append((no_val, topic_val, status_val))
    print(f"{sheet_name}: {len(topics)} topics loaded from Excel. First: {topics[0] if topics else None}, Last: {topics[-1] if topics else None}")

# Check topics-0251-0500.json
try:
    with open('topic-drafts/topics-0251-0500.json', 'r', encoding='utf-8') as f:
        d2 = json.load(f)
    print(f"topics-0251-0500.json: {len(d2)} topics")
    sample_intro = d2[0]['page']['mentor_intro']
    print(f"Sample intro #251: {sample_intro[:100]}...")
except Exception as e:
    print(f"topics-0251-0500.json error: {e}")
