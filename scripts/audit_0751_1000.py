import sqlite3
import json
import glob
import os
import openpyxl

conn = sqlite3.connect('database/app.db')
cur = conn.cursor()

print("=== SQLite Database topic_mentor_drafts by source_batch ===")
batches_db = cur.execute("SELECT source_batch, status, count(*) FROM topic_mentor_drafts GROUP BY source_batch, status").fetchall()
for r in batches_db:
    print(f"  source_batch: {r[0]}, status: {r[1]}, count: {r[2]}")

print("\n=== Scanning topic-drafts/topics-0751-1000.batch*.json ===")
batch_files = sorted(glob.glob('topic-drafts/topics-0751-1000.batch*.json'), key=lambda x: int(x.split('batch')[1].split('.')[0]))
print(f"Found {len(batch_files)} batch files: {batch_files[0]} to {batch_files[-1]}")

all_drafts = []
for bf in batch_files:
    with open(bf, 'r', encoding='utf-8') as f:
        d = json.load(f)
        all_drafts.extend(d)

print(f"Total drafts loaded from batches 1-22: {len(all_drafts)}")
numbers = [d['no'] for d in all_drafts]
print(f"Topic number range: {min(numbers)} to {max(numbers)}")

# Inspect quality metrics
missing_intro = [d['no'] for d in all_drafts if not d.get('page', {}).get('mentor_intro')]
missing_tp = [d['no'] for d in all_drafts if len(d.get('page', {}).get('teaching_points', [])) < 3]
missing_sources = [d['no'] for d in all_drafts if len(d.get('page', {}).get('sources', [])) == 0]
missing_mcq = [d['no'] for d in all_drafts if len(d.get('page', {}).get('mcq_angles', [])) == 0]
missing_hooks = [d['no'] for d in all_drafts if len(d.get('page', {}).get('case_hooks', [])) == 0]

print(f"Missing mentor_intro: {len(missing_intro)}")
print(f"Teaching points < 3: {len(missing_tp)}")
print(f"Missing sources: {len(missing_sources)}")
print(f"Missing mcq_angles: {len(missing_mcq)}")
print(f"Missing case_hooks: {len(missing_hooks)}")

# Check publication status
statuses = set(d.get('status') for d in all_drafts)
print(f"Draft statuses in JSON: {statuses}")

# Check guidelines vs reviews vs trials
designs = {}
for d in all_drafts:
    for s in d.get('page', {}).get('sources', []):
        ds = s.get('design', 'unknown')
        designs[ds] = designs.get(ds, 0) + 1
print(f"Source designs across all 220 topics: {designs}")

# Check master sheet Signal_MD_All_Topics_250_Per_Sheet.xlsx Sheet 4
wb = openpyxl.load_workbook('Signal_MD_All_Topics_250_Per_Sheet.xlsx', data_only=True)
sheet_names = wb.sheetnames
print(f"Workbook sheets: {sheet_names}")
sheet_0751 = None
for s in sheet_names:
    if '0751' in s or '751' in s:
        sheet_0751 = s
        break

if sheet_0751:
    ws = wb[sheet_0751]
    print(f"Master Sheet '{sheet_0751}': max_row = {ws.max_row}")
    col_k_vals = {}
    for r in range(7, min(30, ws.max_row + 1)):
        no = ws.cell(row=r, column=1).value
        top = ws.cell(row=r, column=3).value
        pub = ws.cell(row=r, column=11).value
        ev = ws.cell(row=r, column=8).value
        if r < 12:
            print(f"  Row {r}: No.={no}, Topic='{top}', EvidenceCount={ev}, Published={pub}")
