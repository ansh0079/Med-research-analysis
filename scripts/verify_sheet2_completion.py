import json
import sqlite3
import openpyxl

conn = sqlite3.connect('database/app.db')
cur = conn.cursor()
count_db = cur.execute("SELECT COUNT(*) FROM topic_mentor_drafts WHERE source_batch = 'topics-0251-0500' AND status = 'published'").fetchone()[0]
count_tk = cur.execute("SELECT COUNT(*) FROM topic_knowledge WHERE status = 'human_reviewed'").fetchone()[0]
conn.close()

with open('topic-drafts/topics-0251-0500.json', 'r', encoding='utf-8') as f:
    drafts = json.load(f)
count_json = sum(1 for d in drafts if d.get('status') == 'published')
boilerplate_json = sum(1 for d in drafts if 'crucial clinical entity' in d.get('page', {}).get('mentor_intro', '').lower())

wb_master = openpyxl.load_workbook('Signal_MD_All_Topics_250_Per_Sheet.xlsx', read_only=True)
ws_master = wb_master['Topics 0251-0500']
rows = list(ws_master.iter_rows(values_only=True))[6:]
count_master = sum(1 for r in rows if len(r) > 10 and str(r[10]).lower() == 'published')

wb_rev = openpyxl.load_workbook('topic-drafts/topics-0251-0500-REVIEW.xlsx', read_only=True)
ws_rev = wb_rev.active
rows_rev = list(ws_rev.iter_rows(values_only=True))[1:]
count_rev = sum(1 for r in rows_rev if len(r) > 4 and str(r[4]).lower() == 'published')

print(f"Sheet 2 Verification Summary:")
print(f"1. DB topic_mentor_drafts published: {count_db} / 250")
print(f"2. DB topic_knowledge human_reviewed: {count_tk}")
print(f"3. JSON published drafts: {count_json} / 250 (Boilerplate count: {boilerplate_json})")
print(f"4. Master Workbook published rows: {count_master} / 250")
print(f"5. Review Excel published rows: {count_rev} / 250")
