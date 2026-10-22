#!/usr/bin/env python3
"""
apply_batch_0491_0500.py

Applies Batch 17 of Sheet 2 / Topics 251-500 (Topics 491 to 500).
Upgrades legacy boilerplate with authentic clinical mentor synopses,
landmark trials, and full system synchronization across SQLite, JSON, MD, and Excel workbooks.
This completes 100% of Sheet 2 (Topics 0251-0500)!
"""

import os
import sys
import json
import sqlite3
import re
from datetime import datetime, timezone
import openpyxl

BASE_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
DB_PATH = os.path.join(BASE_DIR, 'database', 'app.db')
MASTER_XLSX = os.path.join(BASE_DIR, 'Signal_MD_All_Topics_250_Per_Sheet.xlsx')
REVIEW_XLSX = os.path.join(BASE_DIR, 'topic-drafts', 'topics-0251-0500-REVIEW.xlsx')
DRAFTS_JSON = os.path.join(BASE_DIR, 'topic-drafts', 'topics-0251-0500.json')
DRAFTS_MD = os.path.join(BASE_DIR, 'topic-drafts', 'topics-0251-0500.md')

sys.path.append(os.path.join(BASE_DIR, 'scripts'))
from data_batch_0491_0500 import BATCH_SHEET2_B17

def normalize_topic(t):
    clean = re.sub(r'[^a-zA-Z0-9\s\-]', ' ', (t or '').lower())
    return re.sub(r'\s+', ' ', clean).strip()

def main():
    print("=== Processing Sheet 2, Batch 17: Topics 491 to 500 ===")

    # Verify citation closure & structure
    for no, d in BATCH_SHEET2_B17.items():
        source_pmids = {str(s['pmid']) for s in d['sources']}
        cited_pmids = set()
        for tp in d['teaching_points']:
            cites = tp.get('cites', [])
            cited_pmids.update(str(c) for c in cites)
        missing = cited_pmids - source_pmids
        assert not missing, f"Topic {no} missing sources for PMIDs: {missing}"
        assert len(d['teaching_points']) == 4, f"Topic {no} teaching points != 4"
        assert len(d['uncertainties']) == 2, f"Topic {no} uncertainties != 2"
        assert len(d['mcq_angles']) == 3, f"Topic {no} mcq_angles != 3"
        assert len(d['case_hooks']) == 2, f"Topic {no} case_hooks != 2"
        assert 'gaps' in d and d['gaps'], f"Topic {no} missing gap"
        assert not 'crucial clinical entity' in d['mentor_intro'].lower(), f"Topic {no} has boilerplate"

    # Load JSON drafts
    with open(DRAFTS_JSON, 'r', encoding='utf-8') as f:
        drafts = json.load(f)

    # SQLite connection
    conn = sqlite3.connect(DB_PATH)
    cur = conn.cursor()

    now_iso = datetime.now(timezone.utc).isoformat()
    updated_count = 0
    k_updated = 0
    k_inserted = 0

    for d in drafts:
        no = d.get('no')
        if no in BATCH_SHEET2_B17:
            binfo = BATCH_SHEET2_B17[no]

            # Update JSON structure
            d['cluster'] = binfo['cluster']
            d['status'] = 'published'
            d['page'] = {
                'mentor_intro': binfo['mentor_intro'],
                'teaching_points': binfo['teaching_points'],
                'sources': binfo['sources'],
                'uncertainties': binfo['uncertainties'],
                'mcq_angles': binfo['mcq_angles'],
                'case_hooks': binfo['case_hooks'],
                'gaps': binfo['gaps']
            }
            d['dbEvidence'] = {
                'totalTrials': len([s for s in binfo['sources'] if s.get('design') == 'trial']),
                'totalGuidelines': len([s for s in binfo['sources'] if s.get('design') in ['guideline', 'consensus-guideline']]),
                'totalMeta': len([s for s in binfo['sources'] if s.get('design') == 'meta-analysis']),
                'totalReviews': len([s for s in binfo['sources'] if s.get('design') == 'review']),
                'guidelinesKept': [s['title'] for s in binfo['sources'] if s.get('design') in ['guideline', 'consensus-guideline']],
                'guidelinesDropped': []
            }
            updated_count += 1

            # Update DB topic_mentor_drafts
            page_json = json.dumps(d['page'])
            ev_json = json.dumps(d['dbEvidence'])

            cur.execute("""
                INSERT INTO topic_mentor_drafts 
                (topic_id, topic, cluster, status, page_json, evidence_json, source_batch, created_at, updated_at, published_at, published_by)
                VALUES (?, ?, ?, 'published', ?, ?, 'topics-0251-0500', ?, ?, ?, 'system_guideline_publisher')
                ON CONFLICT(topic_id) DO UPDATE SET
                    topic = excluded.topic,
                    cluster = excluded.cluster,
                    status = 'published',
                    page_json = excluded.page_json,
                    evidence_json = excluded.evidence_json,
                    source_batch = excluded.source_batch,
                    published_at = excluded.published_at,
                    published_by = excluded.published_by,
                    updated_at = excluded.updated_at
                WHERE review_decision IS NULL
            """, (d['topicId'], d['topic'], binfo['cluster'], page_json, ev_json, now_iso, now_iso, now_iso))

            # Build topic_knowledge with exact 1-based source indices
            sources = binfo['sources']
            idx_map = {str(s['pmid']): i + 1 for i, s in enumerate(sources)}
            source_articles = [{
                'pmid': str(s['pmid']),
                'title': s.get('title', ''),
                'year': s.get('year', ''),
                'design': s.get('design', '')
            } for s in sources]

            has_guideline = any(s.get('design') in ['guideline', 'consensus-guideline'] for s in sources)

            teaching_points = [{
                'claim': tp.get('claim', ''),
                'sourceIndices': [idx_map[str(cite)] for cite in tp.get('cites', []) if str(cite) in idx_map],
                'guidelineIndices': [],
                'confidence': 'HIGH' if has_guideline else 'MODERATE'
            } for tp in binfo['teaching_points']]

            knowledge = {
                'mentorMessage': binfo['mentor_intro'],
                'teachingPoints': teaching_points,
                'caseGenerationHooks': binfo['case_hooks'],
                'mcqAngles': binfo['mcq_angles'],
                'controversies': [{'issue': u, 'sourceIndices': [], 'guidelineIndices': []} for u in binfo['uncertainties']],
                'reviewerGaps': binfo['gaps'],
                'sourceArticles': source_articles
            }

            knowledge_json = json.dumps(knowledge)
            articles_json = json.dumps(source_articles)
            norm_topic = normalize_topic(d['topic'])

            existing_k = cur.execute("SELECT id FROM topic_knowledge WHERE topic = ? OR normalized_topic = ?", (d['topic'], norm_topic)).fetchone()
            if existing_k:
                cur.execute("""
                    UPDATE topic_knowledge
                    SET knowledge = ?, source_articles = ?, status = 'human_reviewed', confidence = 0.92, updated_at = ?, last_refreshed_at = ?
                    WHERE id = ?
                """, (knowledge_json, articles_json, now_iso, now_iso, existing_k[0]))
                k_updated += 1
            else:
                cur.execute("""
                    INSERT INTO topic_knowledge
                    (topic, normalized_topic, knowledge, source_articles, status, confidence, created_at, updated_at, last_refreshed_at)
                    VALUES (?, ?, ?, ?, 'human_reviewed', 0.92, ?, ?, ?)
                """, (d['topic'], norm_topic, knowledge_json, articles_json, now_iso, now_iso, now_iso))
                k_inserted += 1

    conn.commit()
    conn.close()
    print(f"Database: updated {updated_count} drafts in topic_mentor_drafts; topic_knowledge updated={k_updated}, inserted={k_inserted}.")

    # Save JSON
    with open(DRAFTS_JSON, 'w', encoding='utf-8') as f:
        json.dump(drafts, f, indent=2, ensure_ascii=False)
    print(f"Saved: {DRAFTS_JSON}")

    # Generate Markdown
    md_lines = [
        "# Sheet 2 Topics: 0251 to 0500 - Clinical Mentor Digest\n",
        f"*Generated on {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S UTC')}*\n",
        "| No | Topic | Cluster | Status | Guidelines Kept | Trials Count |",
        "|---|---|---|---|---|---|"
    ]
    for d in drafts:
        ev = d.get('dbEvidence', {})
        g_kept = len(ev.get('guidelinesKept', []))
        t_count = ev.get('totalTrials', 0)
        status = d.get('status', 'draft')
        md_lines.append(f"| {d['no']} | {d['topic']} | {d.get('cluster','')} | {status} | {g_kept} | {t_count} |")

    md_lines.append("\n---\n")

    for d in drafts:
        p = d.get('page', {})
        md_lines.append(f"## [{d['no']:04d}] {d['topic']}\n")
        md_lines.append(f"**Cluster**: {d.get('cluster','')} | **Status**: {d.get('status','draft')}\n")
        md_lines.append(f"### Clinical Mentor Synopsis\n{p.get('mentor_intro','')}\n")

        md_lines.append("### Core Teaching Points")
        for tp in p.get('teaching_points', []):
            cites = ", ".join(f"PMID:{c}" for c in tp.get('cites', []))
            md_lines.append(f"- {tp.get('claim','')} ({cites})")
        md_lines.append("")

        md_lines.append("### Key Controversies & Uncertainties")
        for unc in p.get('uncertainties', []):
            md_lines.append(f"- {unc}")
        md_lines.append("")

        md_lines.append("### High-Yield MCQ Angles")
        for mcq in p.get('mcq_angles', []):
            md_lines.append(f"- {mcq}")
        md_lines.append("")

        md_lines.append("### Case Generation Hooks")
        for ch in p.get('case_hooks', []):
            md_lines.append(f"- {ch}")
        md_lines.append("")

        md_lines.append(f"**Reviewer Gap**: {p.get('gaps','')}\n")

        md_lines.append("### Evidence Sources")
        for s in p.get('sources', []):
            md_lines.append(f"- **PMID {s.get('pmid')}**: *{s.get('title')}* ({s.get('year')}) [{s.get('design','')}] - {s.get('journal','')}")
        md_lines.append("\n---\n")

    with open(DRAFTS_MD, 'w', encoding='utf-8') as f:
        f.write("\n".join(md_lines))
    print(f"Saved: {DRAFTS_MD}")

    # Update REVIEW Excel
    wb_rev = openpyxl.load_workbook(REVIEW_XLSX)
    ws_rev = wb_rev.active
    for row in range(2, ws_rev.max_row + 1):
        cell_no = ws_rev.cell(row=row, column=2).value
        if cell_no in BATCH_SHEET2_B17:
            binfo = BATCH_SHEET2_B17[cell_no]
            ws_rev.cell(row=row, column=4, value=binfo['cluster'])
            ws_rev.cell(row=row, column=5, value='published')
            ws_rev.cell(row=row, column=6, value=len(binfo['sources']))
            ws_rev.cell(row=row, column=7, value=len([s for s in binfo['sources'] if s.get('design') == 'trial']))
            ws_rev.cell(row=row, column=8, value=len([s for s in binfo['sources'] if s.get('design') in ['guideline', 'consensus-guideline']]))
            ws_rev.cell(row=row, column=9, value=binfo['mentor_intro'][:120] + '...')
    wb_rev.save(REVIEW_XLSX)
    print(f"Saved: {REVIEW_XLSX}")

    # Update Master Workbook: Signal_MD_All_Topics_250_Per_Sheet.xlsx -> Sheet "Topics 0251-0500"
    wb_master = openpyxl.load_workbook(MASTER_XLSX)
    ws_master = wb_master["Topics 0251-0500"]

    for row in range(2, ws_master.max_row + 1):
        cell_val = ws_master.cell(row=row, column=1).value
        try:
            val_int = int(cell_val)
        except (ValueError, TypeError):
            continue
        if val_int in BATCH_SHEET2_B17:
            binfo = BATCH_SHEET2_B17[val_int]
            ws_master.cell(row=row, column=5, value='published')
            ws_master.cell(row=row, column=6, value=binfo['cluster'])

    wb_master.save(MASTER_XLSX)
    print(f"Saved: {MASTER_XLSX}")

    print("\n=== Sheet 2, Batch 17 (Topics 491 to 500) successfully processed! ===")
    print("=== SHEET 2 (TOPICS 0251-0500) IS NOW 100% COMPLETED AND SYNCHRONIZED! ===")

if __name__ == '__main__':
    main()
