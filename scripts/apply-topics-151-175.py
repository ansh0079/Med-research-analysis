#!/usr/bin/env python3
"""Adopt Topics 151-175 as the next curated/flagship working set."""
from __future__ import annotations

import json
import shutil
from pathlib import Path

import openpyxl

REPO = Path(r"c:\Users\ansh0\OneDrive\Documents\medical research analysis")
XLSX_SRC = Path(r"C:\Users\ansh0\Downloads\Topics_151_to_175_Curated_Literature.xlsx")
OUT_DIR = REPO / "outputs" / "topics-151-175-curated-literature"
CORPUS_PATH = REPO / "data" / "curated-literature-corpus.json"
FLAGSHIP_PATH = REPO / "server" / "config" / "flagshipTopics.json"

# Europe PMC TITLE/AUTH-verified PMIDs only. Unmatched titles stay without pmid.
TITLE_PMIDS = {
    "idsa/aan/acr guidelines for the prevention, diagnosis, and treatment of lyme": ("33558404", 2021),
    "klempner et al": ("11450676", 2001),
    "two controlled trials of antibiotic treatment": ("11450676", 2001),
    "lyme borreliosis (lancet)": ("21903253", 2012),
    "endocrine society clinical practice guideline for testosterone": ("29562364", 2018),
    "traverse trial": ("37733317", 2023),
    "the testosterone trials": ("26886521", 2016),
    "ttrials": ("26886521", 2016),
    "espen guidelines on definitions and terminology": ("27642056", 2017),
    "effort trial": ("31030981", 2019),
    "aha/acc/hfsa guideline for the management of heart failure": ("35378257", 2022),
    "paradigm-hf": ("25176015", 2014),
    "dapa-hf": ("31535829", 2019),
    "emperor-reduced": ("32865377", 2020),
    "rales": ("10471456", 1999),
    "acp clinical practice guideline for noninvasive treatments": ("28192789", 2017),
    "space trial": ("29509867", 2018),
    "nams": ("35797481", 2022),
    "hormone therapy for preventing cardiovascular disease": ("25754617", 2015),
    "women's health initiative": ("12117397", 2002),
    "whi trials": ("12117397", 2002),
    "elite trial": ("27028912", 2016),
    "keeps (kronos": ("32880220", 2021),
    "ukpds": ("9742977", 1998),
    "ada standards of medical care in diabetes": ("36507646", 2023),
    "tear trial": ("22508468", 2012),
    "best study": ("16258899", 2005),
    "levin et al": ("14759425", 2004),
    "methylene blue reduces mortality": ("14759425", 2004),
    "emphasis-hf": ("21073363", 2011),
    "topcat": ("24716680", 2014),
    "mifemiso": ("32853559", 2020),
    "coapt trial": ("30280640", 2018),
    "mitra-fr": ("30145927", 2018),
    "maia trial": ("31141632", 2019),
    "alcyone": ("29231133", 2018),
    "cassiopeia": ("31171419", 2019),
    "swog s0777": ("28017406", 2017),
    "regain trial": ("29066163", 2017),
    "adapt trial": ("34146511", 2021),
    "comfort-i": ("22375971", 2012),
    "comfort-ii": ("22375970", 2012),
    "aap clinical practice guideline revision": ("35927462", 2022),
    "prevent (eculizumab)": ("31050279", 2019),
    "n-momentum": ("31495497", 2019),
    "sakurastar": ("31774956", 2019),
    "checkmate 067": ("31562797", 2019),
    "visual i": ("27602665", 2016),
    "visual ii": ("27542302", 2016),
    "ers/ats official clinical practice guideline: noninvasive ventilation": ("28860265", 2017),
    "brochard et al": ("7651472", 1995),
    "florali": ("25981908", 2015),
}

NEW_FLAGSHIP = [
    {
        "topic": "Malnutrition screening: MUST, NRS-2002, and nutritional support",
        "block": "Gastroenterology",
        "priority": "high",
        "aliases": [
            "EFFORT trial",
            "MUST malnutrition",
            "NRS-2002",
            "ESPEN clinical nutrition",
        ],
        "landmarkPmids": ["31030981", "27642056"],
        "guidelineQueries": [
            "ESPEN ASPEN malnutrition screening MUST NRS-2002 nutritional support guideline",
        ],
        "searchQueries": [
            "EFFORT individualised nutritional support medical inpatients Lancet Schuetz",
            "ESPEN definitions terminology clinical nutrition MUST NRS-2002",
        ],
        "requiredStudyTypes": ["randomized controlled trial", "guideline"],
    },
    {
        "topic": "management of HF with Reduced ejection fraction",
        "block": "Cardiology",
        "priority": "high",
        "aliases": [
            "PARADIGM-HF",
            "DAPA-HF",
            "EMPEROR-Reduced",
            "GDMT HFrEF",
        ],
        "landmarkPmids": ["25176015", "31535829", "32865377", "10471456"],
        "guidelineQueries": [
            "AHA ACC HFSA heart failure reduced ejection fraction GDMT guideline",
        ],
        "searchQueries": [
            "PARADIGM-HF sacubitril valsartan versus enalapril HFrEF NEJM",
            "DAPA-HF EMPEROR-Reduced SGLT2 inhibitor heart failure reduced ejection fraction",
        ],
        "requiredStudyTypes": ["randomized controlled trial", "guideline"],
    },
    {
        "topic": "Mechanical low back pain: assessment, red flags and management",
        "block": "Rheumatology",
        "priority": "medium",
        "aliases": [
            "SPACE trial opioids",
            "ACP low back pain",
            "nonspecific low back pain",
            "red flags back pain",
        ],
        "landmarkPmids": ["28192789", "29509867"],
        "guidelineQueries": [
            "ACP noninvasive treatments acute subacute chronic low back pain guideline",
        ],
        "searchQueries": [
            "SPACE opioid versus nonopioid medications chronic back pain JAMA Krebs",
            "ACP clinical practice guideline noninvasive treatments low back pain Qaseem",
        ],
        "requiredStudyTypes": ["randomized controlled trial", "guideline"],
    },
    {
        "topic": "Menopause and hormone replacement therapy: evidence and guidance",
        "block": "Endocrinology",
        "priority": "medium",
        "aliases": [
            "NICE menopause",
            "KEEPS",
            "WHI long-term follow-up",
            "HRT breast cancer risk",
        ],
        "landmarkPmids": ["12117397", "32880220", "35797481"],
        "guidelineQueries": [
            "NICE menopause diagnosis management hormone replacement therapy guideline",
        ],
        "searchQueries": [
            "Women's Health Initiative estrogen plus progestin postmenopausal women JAMA",
            "Kronos Early Estrogen Prevention Study KEEPS hormone therapy",
        ],
        "requiredStudyTypes": ["randomized controlled trial", "guideline"],
    },
    {
        "topic": "Mental capacity, consent and best-interest decisions",
        "block": "Psychiatry",
        "priority": "medium",
        "aliases": [
            "MacCAT-T",
            "Mental Capacity Act",
            "best interests decision",
            "GMC consent",
        ],
        "landmarkPmids": [],
        "guidelineQueries": [
            "GMC BMA consent assessing mental capacity best interests guideline",
        ],
        "searchQueries": [
            "MacArthur Competence Assessment Tool for Treatment MacCAT-T validation",
            "GMC consent mental capacity best interests clinical practice",
        ],
        "requiredStudyTypes": ["guideline", "review"],
    },
    {
        "topic": "metformin diabetes",
        "block": "Endocrinology",
        "priority": "high",
        "aliases": [
            "UKPDS metformin",
            "UK Prospective Diabetes Study",
            "first-line metformin T2DM",
        ],
        "landmarkPmids": ["9742977"],
        "guidelineQueries": [
            "ADA EASD metformin type 2 diabetes first-line pharmacotherapy guideline",
        ],
        "searchQueries": [
            "UKPDS intensive blood-glucose control metformin overweight type 2 diabetes Lancet",
            "ADA Standards of Care metformin type 2 diabetes cardiovascular",
        ],
        "requiredStudyTypes": ["randomized controlled trial", "guideline"],
    },
    {
        "topic": "methotrexate early rheumatoid arthritis combination therapy TEAR",
        "block": "Rheumatology",
        "priority": "high",
        "aliases": [
            "TEAR trial",
            "BeSt study",
            "triple therapy RA",
            "early aggressive rheumatoid arthritis",
        ],
        "landmarkPmids": ["22508468", "16258899"],
        "guidelineQueries": [
            "ACR EULAR early rheumatoid arthritis methotrexate combination DMARD guideline",
        ],
        "searchQueries": [
            "TEAR oral triple therapy versus etanercept methotrexate early aggressive RA",
            "BeSt four treatment strategies early rheumatoid arthritis radiographic outcomes",
        ],
        "requiredStudyTypes": ["randomized controlled trial", "guideline"],
    },
    {
        "topic": "methylene blue in vasoplegic shock",
        "block": "Critical Care",
        "priority": "medium",
        "aliases": [
            "metylene blue in vasoplegic shock",
            "vasoplegic syndrome methylene blue",
            "post-CPB vasoplegia",
            "nitric oxide synthase inhibitor shock",
        ],
        "landmarkPmids": ["14759425"],
        "guidelineQueries": [
            "Surviving Sepsis Campaign methylene blue vasoplegic shock rescue therapy",
        ],
        "searchQueries": [
            "Levin methylene blue vasoplegic syndrome after cardiac surgery mortality",
            "methylene blue catecholamine-refractory vasoplegia cardiopulmonary bypass",
        ],
        "requiredStudyTypes": ["randomized controlled trial", "guideline"],
    },
    {
        "topic": "Miscarriage: diagnosis, classification and management",
        "block": "Obstetrics",
        "priority": "high",
        "aliases": [
            "MifeMiso",
            "early pregnancy loss",
            "missed miscarriage",
            "ACOG early pregnancy loss",
        ],
        "landmarkPmids": ["32853559"],
        "guidelineQueries": [
            "ACOG RCOG early pregnancy loss miscarriage mifepristone misoprostol guideline",
        ],
        "searchQueries": [
            "MifeMiso mifepristone plus misoprostol versus misoprostol missed miscarriage Lancet",
            "medical versus surgical management early pregnancy loss",
        ],
        "requiredStudyTypes": ["randomized controlled trial", "guideline"],
    },
    {
        "topic": "Neonatal jaundice: assessment, phototherapy and exchange transfusion",
        "block": "Haematology",
        "priority": "medium",
        "aliases": [
            "AAP hyperbilirubinemia",
            "neonatal phototherapy",
            "exchange transfusion jaundice",
            "kernicterus prevention",
        ],
        "landmarkPmids": ["35927462"],
        "guidelineQueries": [
            "AAP hyperbilirubinemia newborn 35 weeks phototherapy exchange transfusion guideline",
        ],
        "searchQueries": [
            "AAP 2022 clinical practice guideline hyperbilirubinemia newborn phototherapy",
            "NICHD aggressive versus conservative phototherapy preterm infants",
        ],
        "requiredStudyTypes": ["guideline", "review"],
    },
    {
        "topic": "nivolumab ipilimumab advanced melanoma overall survival",
        "block": "Oncology",
        "priority": "high",
        "aliases": [
            "CheckMate 067",
            "nivolumab plus ipilimumab melanoma",
            "combined checkpoint inhibition melanoma",
        ],
        "landmarkPmids": ["31562797", "25891173"],
        "guidelineQueries": [
            "ASCO NCCN cutaneous melanoma nivolumab ipilimumab immunotherapy guideline",
        ],
        "searchQueries": [
            "CheckMate 067 five-year survival nivolumab ipilimumab advanced melanoma NEJM",
            "combined nivolumab ipilimumab versus monotherapy untreated melanoma",
        ],
        "requiredStudyTypes": ["randomized controlled trial", "guideline"],
    },
    {
        "topic": "Opioid Rotation and Equianalgesic Dosing",
        "block": "Oncology",
        "priority": "medium",
        "aliases": [
            "opioid switching",
            "equianalgesic table",
            "incomplete cross-tolerance",
            "NCCN adult cancer pain",
        ],
        "landmarkPmids": [],
        "guidelineQueries": [
            "NCCN adult cancer pain opioid rotation equianalgesic dosing guideline",
        ],
        "searchQueries": [
            "Mercadante opioid switching incomplete cross-tolerance cancer pain",
            "opioid rotation equianalgesic tables cancer pain",
        ],
        "requiredStudyTypes": ["guideline", "review"],
    },
]


def lookup_pmid(title: str) -> tuple[str, int] | None:
    key = title.lower()
    for needle, value in TITLE_PMIDS.items():
        if needle in key:
            return value
    return None


def load_excel_topics() -> list[str]:
    wb = openpyxl.load_workbook(XLSX_SRC, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    topics = [str(r[0]).strip() for r in rows[1:] if r and r[0]]
    wb.close()
    return topics


def write_outputs(topics: list[str]) -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(XLSX_SRC, OUT_DIR / XLSX_SRC.name)
    index = OUT_DIR / "topics-index.txt"
    lines = [f"{i}. {t}" for i, t in enumerate(topics, 151)]
    index.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"Wrote {OUT_DIR}")


def enrich_corpus(topics: list[str]) -> None:
    corpus = json.loads(CORPUS_PATH.read_text(encoding="utf-8"))
    wanted = {t.lower() for t in topics}
    missing_topics = wanted - {t["topic"].lower() for t in corpus["topics"]}
    if missing_topics:
        raise SystemExit(f"Corpus missing topics: {sorted(missing_topics)}")
    enriched = 0
    for topic in corpus["topics"]:
        if topic["topic"].lower() not in wanted:
            continue
        for doc in topic["documents"]:
            if doc.get("pmid"):
                continue
            found = lookup_pmid(doc.get("title") or "")
            if not found:
                print(f"  NO PMID <- {(doc.get('title') or '')[:90]}")
                continue
            pmid, year = found
            doc["pmid"] = pmid
            doc["url"] = f"https://pubmed.ncbi.nlm.nih.gov/{pmid}/"
            if year and not doc.get("year"):
                doc["year"] = year
            enriched += 1
            print(f"  PMID {pmid} <- {doc['title'][:90]}")
    CORPUS_PATH.write_text(json.dumps(corpus, indent=2) + "\n", encoding="utf-8")
    print(f"Corpus PMID enrichments: {enriched}")


def add_alias(topic_row: dict, alias: str) -> None:
    aliases = topic_row.setdefault("aliases", [])
    if alias.lower() not in {a.lower() for a in aliases}:
        aliases.append(alias)


def add_pmid(topic_row: dict, pmid: str) -> None:
    pmids = topic_row.setdefault("landmarkPmids", [])
    if pmid not in pmids:
        pmids.append(pmid)


def update_flagship() -> None:
    config = json.loads(FLAGSHIP_PATH.read_text(encoding="utf-8"))
    by_name = {t["topic"].lower(): t for t in config["topics"]}

    lyme = by_name.get("lyme disease diagnosis and treatment")
    if lyme:
        add_alias(lyme, "Klempner persistent Lyme symptoms")
        add_pmid(lyme, "11450676")
        add_pmid(lyme, "33558404")

    hypo = by_name.get("male hypogonadism testosterone replacement")
    if hypo:
        add_alias(hypo, "TTrials")
        add_pmid(hypo, "37733317")
        add_pmid(hypo, "26886521")

    mht = by_name.get("menopausal hormone therapy risks and benefits")
    if mht:
        add_alias(mht, "Menopause and hormone replacement therapy: evidence and guidance")
        add_alias(mht, "ELITE trial")
        add_alias(mht, "NAMS 2022")
        add_pmid(mht, "27028912")
        add_pmid(mht, "35797481")

    hfref = by_name.get("heart failure with reduced ejection fraction")
    if hfref:
        add_alias(hfref, "management of HF with Reduced ejection fraction")
        add_pmid(hfref, "25176015")
        add_pmid(hfref, "31535829")
        add_pmid(hfref, "32865377")

    mra = by_name.get("mineralocorticoid receptor antagonists in heart failure")
    if mra:
        add_alias(mra, "TOPCAT")
        add_pmid(mra, "24716680")

    teer = by_name.get("mitral regurgitation teer mitraclip")
    if teer:
        add_pmid(teer, "30280640")
        add_pmid(teer, "30145927")

    dara = by_name.get("multiple myeloma: daratumumab-based frontline therapy")
    if dara:
        add_alias(dara, "ALCYONE")
        add_alias(dara, "CASSIOPEIA")
        add_pmid(dara, "29231133")
        add_pmid(dara, "31171419")

    nmosd = by_name.get("neuromyelitis optica spectrum disorder biologic therapy")
    if nmosd:
        add_alias(nmosd, "SAkuraStar")
        add_pmid(nmosd, "31774956")

    niv = by_name.get("noninvasive ventilation for acute respiratory failure")
    if niv:
        add_alias(niv, "FLORALI")
        add_pmid(niv, "25981908")
        add_pmid(niv, "28860265")
        add_pmid(niv, "7651472")

    ra = by_name.get("rheumatoid arthritis initial dmard strategy")
    if ra:
        add_alias(ra, "methotrexate early rheumatoid arthritis combination therapy TEAR")
        add_alias(ra, "BeSt study")
        add_pmid(ra, "22508468")

    melanoma = by_name.get("advanced melanoma immunotherapy")
    if melanoma:
        add_alias(melanoma, "nivolumab ipilimumab advanced melanoma overall survival")
        add_pmid(melanoma, "31562797")

    t2dm = by_name.get("type 2 diabetes cardiovascular outcomes")
    if t2dm:
        add_alias(t2dm, "metformin diabetes")
        add_alias(t2dm, "UKPDS")
        add_pmid(t2dm, "9742977")

    existing = {t["topic"].lower() for t in config["topics"]}
    added = 0
    for row in NEW_FLAGSHIP:
        if row["topic"].lower() in existing:
            continue
        config["topics"].append(row)
        existing.add(row["topic"].lower())
        added += 1

    config["version"] = 16
    config["targets"]["flagshipCount"] = len(config["topics"])
    config["description"] = (
        f"Flagship-topic program - {len(config['topics'])} topics spanning evidence + learning surfaces, "
        "including curated literature cohorts Topics 51-75, Topics 76-100, Topics 126-150, and Topics 151-175."
    )
    FLAGSHIP_PATH.write_text(json.dumps(config, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"Flagship topics now {len(config['topics'])} (added {added})")


def main() -> None:
    if not XLSX_SRC.exists():
        raise SystemExit(f"Missing {XLSX_SRC}")
    topics = load_excel_topics()
    if len(topics) != 25:
        raise SystemExit(f"Expected 25 topics, found {len(topics)}")
    write_outputs(topics)
    enrich_corpus(topics)
    update_flagship()


if __name__ == "__main__":
    main()
