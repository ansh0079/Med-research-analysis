import sqlite3
import json
import openpyxl

print("=== FINAL 5-LAYER AUDIT: SHEET 5 & MASTER WORKBOOK ===")

# 1. SQLite verification
conn = sqlite3.connect('database/app.db')
c = conn.cursor()

with open('topic-drafts/topics-1001-1250.json', 'r', encoding='utf-8') as f:
    drafts = json.load(f)

json_pub_250 = [d for d in drafts if d['status'] == 'published' and 1001 <= d['no'] <= 1250]
print(f"1. JSON published in Sheet 5 (1001-1250): {len(json_pub_250)} / 250")

topic_ids = [d['topicId'] for d in json_pub_250]
placeholders = ','.join('?' for _ in topic_ids)
c.execute(f"SELECT COUNT(*) FROM topic_mentor_drafts WHERE status = 'published' AND topic_id IN ({placeholders})", topic_ids)
db_pub = c.fetchone()[0]
print(f"2. DB topic_mentor_drafts published in Sheet 5: {db_pub} / 250")

# 3. Master sheet check
wb = openpyxl.load_workbook('Signal_MD_All_Topics_250_Per_Sheet.xlsx', data_only=True)
ws = wb['Topics 1001-1250']
master_pub = 0
for r in range(7, 257):
    val = ws.cell(row=r, column=11).value
    if val == 'Published':
        master_pub += 1
print(f"3. Master Sheet published in Sheet 5: {master_pub} / 250")

# 4. Review sheet check
wb_rev = openpyxl.load_workbook('topic-drafts/topics-1001-1250-REVIEW.xlsx', data_only=True)
ws_rev = wb_rev.active
rev_pub = 0
for r in range(2, 252):
    val = ws_rev.cell(row=r, column=5).value
    if val == 'Published':
        rev_pub += 1
print(f"4. Review Sheet published in Sheet 5: {rev_pub} / 250")

# 5. Markdown check
with open('topic-drafts/topics-1001-1250.md', 'r', encoding='utf-8') as f:
    md_content = f.read()
md_headers = md_content.count('## ')
print(f"5. Markdown topic headers in Sheet 5: {md_headers} / 250")

print("\n--- MASTER WORKBOOK AUDIT (ALL 8 SHEETS) ---")
total_published = 0
total_topics = 0
for sheet in wb.sheetnames:
    ws_s = wb[sheet]
    s_pub = 0
    s_tot = 0
    for r in range(7, ws_s.max_row + 1):
        num = ws_s.cell(row=r, column=1).value
        if num is not None and isinstance(num, int):
            s_tot += 1
            if ws_s.cell(row=r, column=11).value == 'Published':
                s_pub += 1
    total_published += s_pub
    total_topics += s_tot
    pct = (s_pub / s_tot * 100) if s_tot > 0 else 0
    print(f"  * {sheet}: {s_pub} / {s_tot} ({pct:.1f}%) published")

print(f"\n=======================================================")
print(f"TOTAL CURRICULUM TOPICS: {total_published} / {total_topics} PUBLISHED ({total_published/total_topics*100:.1f}%)")
print(f"=======================================================")

wb.close()
wb_rev.close()
conn.close()
