import os
import openpyxl

base = 'topic-drafts'
files = os.listdir(base)
print("Files in topic-drafts:", files)

wb = openpyxl.load_workbook('Signal_MD_All_Topics_250_Per_Sheet.xlsx', read_only=True)
print("Master Workbook Sheets:", wb.sheetnames)
for sname in wb.sheetnames:
    ws = wb[sname]
    print(f"Sheet {sname}: {ws.max_row} rows")
