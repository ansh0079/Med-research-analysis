#!/usr/bin/env python3
"""Fill Highly_Cited_Literature_Index.xlsx from corpus URLs and local DB status."""
from __future__ import annotations

import json
import sqlite3
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

REPO = Path(r"c:\Users\ansh0\OneDrive\Documents\medical research analysis")
SRC = Path(r"c:\Users\ansh0\Downloads\Highly_Cited_Literature_Index.xlsx")
OUT_DIR = REPO / "outputs" / "highly-cited-literature-index"
OUT_XLSX = OUT_DIR / "Highly_Cited_Literature_Index_with_links.xlsx"
CORPUS_PATH = REPO / "data" / "curated-literature-corpus.json"
DB_PATH = REPO / "database" / "app.db"
ROWS_JSON = OUT_DIR / "index-rows.json"

REVIEW_TYPES = {"review_article"}
META_TYPES = {"systematic_review_meta_analysis", "meta_analysis"}


def key(s: str) -> str:
    return " ".join("".join(ch.lower() if ch.isalnum() else " " for ch in str(s or "")).split())


def url_for(doc: dict) -> str | None:
    if doc.get("url"):
        return str(doc["url"]).strip()
    if doc.get("pmcid"):
        pmc = str(doc["pmcid"]).upper()
        if not pmc.startswith("PMC"):
            pmc = f"PMC{pmc}"
        return f"https://pmc.ncbi.nlm.nih.gov/articles/{pmc}/"
    if doc.get("pmid"):
        return f"https://pubmed.ncbi.nlm.nih.gov/{doc['pmid']}/"
    return None


def pick(docs: list[dict], types: set[str]) -> dict | None:
    ranked = [d for d in docs if d.get("type") in types]
    with_id = [d for d in ranked if url_for(d)]
    return (with_id or ranked or [None])[0]


def local_status(con: sqlite3.Connection, docs: list[dict]) -> str:
    pmids = [str(d["pmid"]) for d in docs if d.get("pmid")]
    pmcids = [str(d["pmcid"]).upper().replace("PMC", "PMC") for d in docs if d.get("pmcid")]
    titles = [str(d["title"]).strip() for d in docs if d.get("title")]
    rows = []
    if pmids:
        q = ",".join("?" * len(pmids))
        rows += con.execute(
            f"SELECT full_text_source, word_count FROM guideline_documents WHERE pmid IN ({q})",
            pmids,
        ).fetchall()
    if pmcids:
        q = ",".join("?" * len(pmcids))
        rows += con.execute(
            f"SELECT full_text_source, word_count FROM guideline_documents WHERE pmcid IN ({q})",
            pmcids,
        ).fetchall()
    if not rows and titles:
        # exact title match only — do not fuzzy-guess
        q = ",".join("?" * len(titles))
        rows += con.execute(
            f"SELECT full_text_source, word_count FROM guideline_documents WHERE title IN ({q})",
            titles,
        ).fetchall()
    if any(src == "jats" and (wc or 0) >= 500 for src, wc in rows):
        return "Local JATS body"
    if any(src == "manual" and (wc or 0) >= 500 for src, wc in rows):
        return "Local HTML/manual body"
    if any(src == "abstract" for src, wc in rows):
        return "Local abstract only"
    if rows:
        return "Local metadata only"
    return "No local body"


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    corpus = json.loads(CORPUS_PATH.read_text(encoding="utf-8"))
    by_topic = {key(t["topic"]): t for t in corpus["topics"]}
    con = sqlite3.connect(DB_PATH)

    wb = openpyxl.load_workbook(SRC)
    ws = wb["Literature Index"]
    ws["F1"] = "Review title"
    ws["G1"] = "Review link status"
    ws["H1"] = "Meta-analysis title"
    ws["I1"] = "Meta-analysis link status"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1E3A8A")
    link_font = Font(color="0563C1", underline="single")
    thin = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB"),
    )
    for col in range(1, 10):
        cell = ws.cell(1, col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    counts = {
        "review_url": 0,
        "review_title_only": 0,
        "review_missing": 0,
        "meta_url": 0,
        "meta_title_only": 0,
        "meta_missing": 0,
        "unmatched_topic": 0,
    }
    exported = []

    for row in range(2, ws.max_row + 1):
        name = ws.cell(row, 2).value
        entry = by_topic.get(key(name or ""))
        if not entry:
            ws.cell(row, 3).value = "Unmatched topic"
            counts["unmatched_topic"] += 1
            exported.append({"topic": name, "reviewUrl": None, "metaUrl": None})
            continue
        docs = entry.get("documents") or []
        review = pick(docs, REVIEW_TYPES)
        meta = pick(docs, META_TYPES)
        review_url = url_for(review) if review else None
        meta_url = url_for(meta) if meta else None
        if review_url:
            counts["review_url"] += 1
            review_status = "pubmed" if "pubmed" in review_url else "url"
        elif review:
            counts["review_title_only"] += 1
            review_status = "title_only"
        else:
            counts["review_missing"] += 1
            review_status = "missing"

        if meta_url:
            counts["meta_url"] += 1
            meta_status = "pubmed" if "pubmed" in meta_url else "url"
        elif meta:
            counts["meta_title_only"] += 1
            meta_status = "title_only"
        else:
            counts["meta_missing"] += 1
            meta_status = "missing"

        ws.cell(row, 3).value = local_status(con, docs)
        review_cell = ws.cell(row, 4)
        review_cell.value = review_url
        if review_url:
            review_cell.hyperlink = review_url
            review_cell.font = link_font
        meta_cell = ws.cell(row, 5)
        meta_cell.value = meta_url
        if meta_url:
            meta_cell.hyperlink = meta_url
            meta_cell.font = link_font
        ws.cell(row, 6).value = (review or {}).get("title")
        ws.cell(row, 7).value = review_status
        ws.cell(row, 8).value = (meta or {}).get("title")
        ws.cell(row, 9).value = meta_status
        for col in range(1, 10):
            c = ws.cell(row, col)
            c.border = thin
            c.alignment = Alignment(wrap_text=True, vertical="center")
        exported.append({
            "id": ws.cell(row, 1).value,
            "topic": name,
            "reviewTitle": (review or {}).get("title"),
            "reviewUrl": review_url,
            "reviewPmid": (review or {}).get("pmid"),
            "reviewPmcid": (review or {}).get("pmcid"),
            "metaTitle": (meta or {}).get("title"),
            "metaUrl": meta_url,
            "metaPmid": (meta or {}).get("pmid"),
            "metaPmcid": (meta or {}).get("pmcid"),
            "localStatus": ws.cell(row, 3).value,
        })

    widths = [10, 48, 22, 45, 45, 50, 18, 50, 18]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.auto_filter.ref = f"A1:I{ws.max_row}"
    ws.freeze_panes = "A2"
    ws.column_dimensions["C"].width = 24

    con.close()
    wb.save(OUT_XLSX)
    wb.save(SRC)
    ROWS_JSON.write_text(json.dumps(exported, indent=2), encoding="utf-8")
    print(json.dumps({"out": str(OUT_XLSX), "rows": len(exported), "counts": counts}, indent=2))


if __name__ == "__main__":
    main()
