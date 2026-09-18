#!/usr/bin/env python3
"""Fill Clinical_Guidelines_Index.xlsx Search Guideline Link from known corpus/DB URLs."""
from __future__ import annotations

import json
import shutil
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

REPO = Path(r"c:\Users\ansh0\OneDrive\Documents\medical research analysis")
SRC = Path(r"c:\Users\ansh0\Downloads\Clinical_Guidelines_Index.xlsx")
OUT_DIR = REPO / "outputs" / "clinical-guidelines-index"
OUT_XLSX = OUT_DIR / "Clinical_Guidelines_Index_with_links.xlsx"
CORPUS_PATH = REPO / "data" / "curated-literature-corpus.json"
DB_PATH = REPO / "database" / "app.db"

# Previously verified public guideline pages (not guessed).
VERIFIED_PAGES = {
    "achalasia endoscopic and surgical therapy": (
        "https://pmc.ncbi.nlm.nih.gov/articles/PMC9896940/",
        "ACG Clinical Guidelines: Diagnosis and Management of Achalasia",
        "verified_oa_page",
    ),
    "acute bacterial meningitis adjunctive dexamethasone": (
        "https://www.who.int/publications/i/item/9789240108042",
        "WHO guidelines on meningitis diagnosis, treatment and care",
        "verified_oa_page",
    ),
    "acute bacterial meningitis empiric therapy and dexamethasone": (
        "https://www.nice.org.uk/guidance/ng240",
        "NICE NG240: Meningitis (bacterial) and meningococcal disease",
        "verified_oa_page",
    ),
    "acute liver failure transplant referral": (
        "https://pubmed.ncbi.nlm.nih.gov/28417882/",
        "EASL Clinical Practical Guidelines on acute fulminant liver failure",
        "pubmed",
    ),
    "acute myeloid leukemia venetoclax azacitidine": (
        "https://ashpublications.org/blood/article/140/12/1345/485817/Diagnosis-and-management-of-AML-in-adults-2022",
        "European LeukemiaNet 2022 AML recommendations",
        "verified_oa_page",
    ),
}


def key(s: str) -> str:
    out = []
    for ch in str(s or "").lower():
        out.append(ch if ch.isalnum() else " ")
    return " ".join("".join(out).split())


def load_jats_cpgs() -> dict[str, dict]:
    try:
        import sqlite3
        con = sqlite3.connect(DB_PATH)
        rows = con.execute(
            """
            SELECT title, pmid, pmcid, word_count
            FROM guideline_documents
            WHERE document_label = 'clinical_practice_guideline'
              AND full_text_source = 'jats'
              AND full_text IS NOT NULL
              AND LENGTH(full_text) >= 2000
              AND title LIKE '%Guideline%'
            """
        ).fetchall()
        con.close()
    except Exception:
        return {}
    by_pmid: dict[str, dict] = {}
    by_title: dict[str, dict] = {}
    for title, pmid, pmcid, words in rows:
        rec = {"title": title, "pmid": pmid, "pmcid": pmcid, "words": words}
        if pmid:
            by_pmid[str(pmid)] = rec
        by_title[key(title)] = rec
    return {"pmid": by_pmid, "title": by_title}


def guideline_docs(topic: dict) -> list[dict]:
    return [d for d in topic.get("documents") or [] if d.get("type") == "clinical_practice_guideline"]


def pick_link(topic_name: str, docs: list[dict], jats: dict) -> tuple[str | None, str | None, str]:
    k = key(topic_name)
    if k in VERIFIED_PAGES:
        url, title, status = VERIFIED_PAGES[k]
        return url, title, status

    for d in docs:
        pmid = str(d.get("pmid") or "")
        title = d.get("title") or ""
        hit = (pmid and jats.get("pmid", {}).get(pmid)) or jats.get("title", {}).get(key(title))
        if hit:
            pmcid = hit.get("pmcid") or d.get("pmcid")
            if pmcid:
                pmc = str(pmcid).upper()
                if not pmc.startswith("PMC"):
                    pmc = f"PMC{pmc}"
                return f"https://pmc.ncbi.nlm.nih.gov/articles/{pmc}/", hit["title"], "full_jats"
            if hit.get("pmid"):
                return f"https://pubmed.ncbi.nlm.nih.gov/{hit['pmid']}/", hit["title"], "full_jats"

    for d in docs:
        url = (d.get("url") or "").strip()
        pmcid = d.get("pmcid")
        pmid = d.get("pmid")
        title = d.get("title")
        if pmcid:
            pmc = str(pmcid).upper()
            if not pmc.startswith("PMC"):
                pmc = f"PMC{pmc}"
            return f"https://pmc.ncbi.nlm.nih.gov/articles/{pmc}/", title, "pmc_record"
        if url:
            status = "pubmed" if "pubmed" in url.lower() else "url"
            return url, title, status
        if pmid:
            return f"https://pubmed.ncbi.nlm.nih.gov/{pmid}/", title, "pubmed"

    if docs:
        return None, docs[0].get("title"), "title_only"
    return None, None, "no_guideline_title"


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SRC, OUT_DIR / "Clinical_Guidelines_Index.source.xlsx")

    corpus = json.loads(CORPUS_PATH.read_text(encoding="utf-8"))
    by_topic = {key(t["topic"]): t for t in corpus["topics"]}
    jats = load_jats_cpgs()

    wb = openpyxl.load_workbook(SRC)
    ws = wb["Guidelines Index"]
    ws["E1"] = "Guideline title"
    ws["F1"] = "Link status"
    ws["G1"] = "Full guideline downloaded"
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="14532D")
    for col in range(1, 8):
        cell = ws.cell(1, col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    link_font = Font(color="0563C1", underline="single")
    thin = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB"),
    )
    counts = {"full_jats": 0, "verified_oa_page": 0, "pmc_record": 0, "pubmed": 0, "url": 0, "title_only": 0, "no_guideline_title": 0, "unmatched_topic": 0}

    for row in range(2, ws.max_row + 1):
        name = ws.cell(row, 2).value
        entry = by_topic.get(key(name or ""))
        if not entry:
            ws.cell(row, 6).value = "unmatched_topic"
            counts["unmatched_topic"] += 1
            continue
        docs = guideline_docs(entry)
        url, title, status = pick_link(name, docs, jats)
        counts[status] = counts.get(status, 0) + 1
        link_cell = ws.cell(row, 4)
        if url:
            link_cell.value = url
            link_cell.hyperlink = url
            link_cell.font = link_font
        else:
            link_cell.value = None
        ws.cell(row, 5).value = title
        ws.cell(row, 6).value = status
        ws.cell(row, 7).value = "yes" if status == "full_jats" else "no"
        for col in range(1, 8):
            c = ws.cell(row, col)
            c.border = thin
            c.alignment = Alignment(wrap_text=True, vertical="center")

    widths = [8, 48, 32, 55, 55, 18, 22]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.auto_filter.ref = f"A1:G{ws.max_row}"
    ws.freeze_panes = "A2"

    wb.save(OUT_XLSX)
    shutil.copy2(OUT_XLSX, SRC)
    print(json.dumps({"out": str(OUT_XLSX), "rows": ws.max_row - 1, "counts": counts}, indent=2))


if __name__ == "__main__":
    main()
