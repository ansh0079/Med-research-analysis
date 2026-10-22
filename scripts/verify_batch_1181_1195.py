import sqlite3
import json
import openpyxl

conn = sqlite3.connect('database/app.db')
c = conn.cursor()

with open('topic-drafts/topics-1001-1250.json', 'r', encoding='utf-8') as f:
    drafts = json.load(f)

json_pub_195 = [d for d in drafts if d['status'] == 'published' and 1001 <= d['no'] <= 1195]
print(f"JSON published in 1001-1195: {len(json_pub_195)} / 195")

all_pub_json = [d for d in drafts if d['status'] == 'published']
print(f"JSON total published in Sheet 5: {len(all_pub_json)} / 250")

topic_ids = [d['topicId'] for d in json_pub_195]
placeholders = ','.join('?' for _ in topic_ids)
c.execute(f"SELECT COUNT(*) FROM topic_mentor_drafts WHERE status = 'published' AND topic_id IN ({placeholders})", topic_ids)
db_pub = c.fetchone()[0]
print(f"DB topic_mentor_drafts published in 1001-1195: {db_pub} / 195")

# Master sheet
wb = openpyxl.load_workbook('Signal_MD_All_Topics_250_Per_Sheet.xlsx', data_only=True)
ws = wb['Topics 1001-1250']
master_pub = 0
for r in range(6, 201):
    val = ws.cell(row=r, column=11).value
    if val == 'Published':
        master_pub += 1
print(f"Master Sheet published in 1001-1195: {master_pub} / 195")

# Review sheet
wb_rev = openpyxl.load_workbook('topic-drafts/topics-1001-1250-REVIEW.xlsx', data_only=True)
ws_rev = wb_rev.active
rev_pub = 0
for r in range(2, 197):
    val = ws_rev.cell(row=r, column=5).value
    if val == 'Published':
        rev_pub += 1
print(f"Review Sheet published in 1001-1195: {rev_pub} / 195")

wb.close()
wb_rev.close()
conn.close()
print("All layer checks completed successfully!")
