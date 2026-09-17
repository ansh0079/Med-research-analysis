#!/usr/bin/env python3
"""Adopt Topics 201-225 as the next curated/flagship working set."""
from __future__ import annotations

import json
import shutil
from pathlib import Path

import openpyxl

REPO = Path(r"c:\Users\ansh0\OneDrive\Documents\medical research analysis")
XLSX_SRC = Path(r"C:\Users\ansh0\Downloads\Topics_201_to_225_Curated_Literature.xlsx")
OUT_DIR = REPO / "outputs" / "topics-201-225-curated-literature"
CORPUS_PATH = REPO / "data" / "curated-literature-corpus.json"
FLAGSHIP_PATH = REPO / "server" / "config" / "flagshipTopics.json"

# Europe PMC TITLE/AUTH-verified PMIDs only. Unmatched titles stay without pmid.
TITLE_PMIDS = {
    "eular/acr recommendations for the management of polymyalgia": ("26359488", 2015),
    "sarilumab for pmr": ("37792612", 2023),
    "baveno vii": ("35120736", 2022),
    "predesci": ("30910320", 2019),
    "consensus guidelines for the management of postoperative nausea": ("32467512", 2020),
    "apfel et al": ("10485781", 1999),
    "hrs expert consensus statement": ("25980576", 2015),
    "exercise training versus beta-blocker": ("21690484", 2011),
    "endocrine society clinical practice guideline on the management of primary aldosteronism": (
        "26934393",
        2016,
    ),
    "pathway-2": ("26414968", 2015),
    "kdigo clinical practice guideline for glomerular diseases": ("34556256", 2021),
    "mentor trial": ("31269364", 2019),
    "madit-ii": ("11907286", 2002),
    "scd-heft": ("15659722", 2005),
    "endocrine society clinical practice guideline on diagnosis and treatment of hyperprolactinemia": (
        "21296991",
        2011,
    ),
    "webster et al": ("7915824", 1994),
    "proseva": ("23688302", 2013),
    "idsa guidelines for the diagnosis and management of prosthetic joint": ("23223583", 2013),
    "zimmerli et al": ("15483283", 2004),
    "poort et al": ("8916933", 1996),
    "aad-npf guidelines of care for the management of psoriasis with biologics": ("30772098", 2019),
    "voyage trials (guselkumab)": ("28057361", 2017),
    "eclipse trial": ("31402114", 2019),
    "foa et al": ("16287395", 2005),
    "va/dod clinical practice guideline for management of posttraumatic": ("32021581", 2018),
    "davidson et al": ("11343529", 2001),
    "fleischner society guidelines for incidental pulmonary nodules": ("28240562", 2017),
    "national lung screening trial": ("21714641", 2011),
    "nelson trial": ("31995683", 2020),
    "baughman et al": ("16840744", 2006),
    "advisory committee on immunization practices": ("18496505", 2008),
    "rival trial": ("21470671", 2011),
    "matrix trial": ("25791214", 2015),
    "aha/asa adult stroke rehabilitation": ("27145936", 2016),
    "avert trial": ("25892679", 2015),
}

NEW_FLAGSHIP = [
    {
        "topic": "Postoperative ileus: prevention and management",
        "block": "Gastroenterology",
        "priority": "medium",
        "aliases": [
            "postoperative ileus",
            "ERAS ileus",
            "gum chewing ileus",
        ],
        "landmarkPmids": [],
        "guidelineQueries": [
            "ERAS enhanced recovery after surgery postoperative ileus prevention guideline",
        ],
        "searchQueries": [
            "gum chewing early enteral nutrition postoperative ileus Cochrane",
            "ERAS colorectal surgery ileus prevention",
        ],
        "requiredStudyTypes": ["guideline", "systematic review"],
    },
    {
        "topic": "Postoperative nausea and vomiting: prevention and treatment",
        "block": "Emergency Medicine",
        "priority": "medium",
        "aliases": [
            "PONV",
            "Apfel score",
            "postoperative antiemetic prophylaxis",
        ],
        "landmarkPmids": ["32467512", "10485781"],
        "guidelineQueries": [
            "consensus guidelines postoperative nausea vomiting Anesthesia Analgesia",
        ],
        "searchQueries": [
            "Apfel simplified risk score postoperative nausea vomiting",
            "multimodal antiemetic prophylaxis PONV randomized",
        ],
        "requiredStudyTypes": ["guideline", "randomized controlled trial"],
    },
    {
        "topic": "Postural Orthostatic Tachycardia Syndrome (POTS): tilt table test criteria, non-pharmacological methods",
        "block": "Cardiology",
        "priority": "medium",
        "aliases": [
            "POTS",
            "postural tachycardia syndrome",
            "HRS POTS consensus",
            "Fu exercise POTS",
        ],
        "landmarkPmids": ["25980576", "21690484"],
        "guidelineQueries": [
            "Heart Rhythm Society expert consensus postural tachycardia syndrome POTS",
        ],
        "searchQueries": [
            "Fu exercise training versus propranolol postural orthostatic tachycardia",
            "HRS expert consensus diagnosis treatment POTS tilt table",
        ],
        "requiredStudyTypes": ["guideline", "randomized controlled trial"],
    },
    {
        "topic": "Pressure ulcers: staging, prevention and wound management",
        "block": "Dermatology",
        "priority": "medium",
        "aliases": [
            "pressure injury",
            "NPUAP EPUAP",
            "support surfaces pressure ulcer",
        ],
        "landmarkPmids": [],
        "guidelineQueries": [
            "NPUAP EPUAP PPPIA international clinical practice guideline pressure injury",
        ],
        "searchQueries": [
            "support surfaces specialized mattresses pressure ulcer prevention Cochrane",
            "hospital repositioning pressure injury prevention",
        ],
        "requiredStudyTypes": ["guideline", "systematic review"],
    },
    {
        "topic": "prone positioning severe acute respiratory distress syndrome",
        "block": "Critical Care",
        "priority": "high",
        "aliases": [
            "PROSEVA",
            "prone positioning severe ARDS",
            "proning ARDS",
        ],
        "landmarkPmids": ["23688302"],
        "guidelineQueries": [
            "ESICM ARDS prone positioning Surviving Sepsis guideline",
        ],
        "searchQueries": [
            "PROSEVA prone positioning severe ARDS NEJM Guerin",
            "duration prone positioning mortality moderate severe ARDS",
        ],
        "requiredStudyTypes": ["randomized controlled trial", "guideline"],
    },
    {
        "topic": "Prothrombin G20210A Mutation",
        "block": "Haematology",
        "priority": "medium",
        "aliases": [
            "prothrombin gene mutation",
            "factor II G20210A",
            "Poort prothrombin",
        ],
        "landmarkPmids": ["8916933"],
        "guidelineQueries": [
            "ASH venous thromboembolism inherited thrombophilia prothrombin G20210A",
        ],
        "searchQueries": [
            "Poort prothrombin G20210A 3 untranslated region Blood 1996",
            "prothrombin G20210A venous thromboembolism risk",
        ],
        "requiredStudyTypes": ["guideline", "review"],
    },
    {
        "topic": "Psoriasis: diagnosis, severity assessment and systemic treatment",
        "block": "Dermatology",
        "priority": "high",
        "aliases": [
            "AAD-NPF psoriasis biologics",
            "VOYAGE guselkumab",
            "ECLIPSE guselkumab secukinumab",
            "plaque psoriasis systemic therapy",
        ],
        "landmarkPmids": ["30772098", "28057361", "31402114"],
        "guidelineQueries": [
            "AAD NPF guidelines care management psoriasis biologics",
        ],
        "searchQueries": [
            "VOYAGE guselkumab versus adalimumab plaque psoriasis",
            "ECLIPSE guselkumab versus secukinumab moderate severe psoriasis",
        ],
        "requiredStudyTypes": ["randomized controlled trial", "guideline"],
    },
    {
        "topic": "Psychosis and acute behavioural disturbance",
        "block": "Psychiatry",
        "priority": "medium",
        "aliases": [
            "acute agitation",
            "Project BETA",
            "NICE NG10 violence aggression",
        ],
        "landmarkPmids": [],
        "guidelineQueries": [
            "NICE NG10 violence aggression short-term management mental health",
        ],
        "searchQueries": [
            "intramuscular antipsychotics benzodiazepines acute agitation comparative",
            "Project BETA evaluation treatment agitation emergency",
        ],
        "requiredStudyTypes": ["guideline", "review"],
    },
    {
        "topic": "ptsd trauma focused psychotherapy and ssri",
        "block": "Psychiatry",
        "priority": "high",
        "aliases": [
            "prolonged exposure PTSD",
            "Foa prolonged exposure",
            "APA PTSD guideline",
        ],
        "landmarkPmids": ["16287395"],
        "guidelineQueries": [
            "APA clinical practice guideline treatment posttraumatic stress disorder",
        ],
        "searchQueries": [
            "Foa prolonged exposure therapy PTSD randomized trial",
            "trauma focused psychotherapy versus SSRI PTSD",
        ],
        "requiredStudyTypes": ["randomized controlled trial", "guideline"],
    },
    {
        "topic": "PTSD trauma-focused psychotherapy and SSRI",
        "block": "Psychiatry",
        "priority": "high",
        "aliases": [
            "VA DoD PTSD guideline",
            "sertraline PTSD",
            "Davidson sertraline PTSD",
        ],
        "landmarkPmids": ["32021581", "11343529"],
        "guidelineQueries": [
            "VA DoD clinical practice guideline management posttraumatic stress disorder",
        ],
        "searchQueries": [
            "Davidson sertraline placebo posttraumatic stress disorder multicenter",
            "SSRI SNRI efficacy PTSD randomized trial",
        ],
        "requiredStudyTypes": ["randomized controlled trial", "guideline"],
    },
    {
        "topic": "pulmonary nodule evaluation and fleischner follow up",
        "block": "Respiratory",
        "priority": "high",
        "aliases": [
            "Fleischner pulmonary nodules",
            "NLST",
            "NELSON",
            "incidental pulmonary nodule",
        ],
        "landmarkPmids": ["28240562", "21714641", "31995683"],
        "guidelineQueries": [
            "Fleischner Society incidental pulmonary nodules CT follow-up guideline",
        ],
        "searchQueries": [
            "NLST reduced lung cancer mortality low-dose CT screening NEJM",
            "NELSON volume CT screening lung cancer mortality randomized",
        ],
        "requiredStudyTypes": ["randomized controlled trial", "guideline"],
    },
    {
        "topic": "Pulmonary sarcoidosis corticosteroid therapy",
        "block": "Respiratory",
        "priority": "medium",
        "aliases": [
            "ATS sarcoidosis treatment",
            "Baughman infliximab sarcoidosis",
            "pulmonary sarcoidosis methotrexate",
        ],
        "landmarkPmids": ["16840744"],
        "guidelineQueries": [
            "ATS clinical practice guideline treatment sarcoidosis corticosteroids",
        ],
        "searchQueries": [
            "Baughman infliximab chronic sarcoidosis pulmonary involvement",
            "corticosteroids versus immunosuppressants pulmonary sarcoidosis",
        ],
        "requiredStudyTypes": ["randomized controlled trial", "guideline"],
    },
    {
        "topic": "Rabies pre- and post-exposure prophylaxis",
        "block": "Infectious Diseases",
        "priority": "high",
        "aliases": [
            "rabies PEP",
            "ACIP rabies",
            "WHO rabies prophylaxis",
        ],
        "landmarkPmids": ["18496505"],
        "guidelineQueries": [
            "ACIP WHO rabies pre-exposure post-exposure prophylaxis immunoglobulin",
        ],
        "searchQueries": [
            "ACIP human rabies prevention United States 2008 recommendations",
            "rabies vaccine immunoglobulin post-exposure regimen efficacy",
        ],
        "requiredStudyTypes": ["guideline", "review"],
    },
    {
        "topic": "Radiation necrosis vs tumour progression: perfusion MRI and PET scan utility",
        "block": "Oncology",
        "priority": "medium",
        "aliases": [
            "radiation necrosis PET",
            "RANO radiation necrosis",
            "amino acid PET glioma",
        ],
        "landmarkPmids": [],
        "guidelineQueries": [
            "RANO response assessment neuro-oncology radiation necrosis versus progression",
        ],
        "searchQueries": [
            "amino acid PET perfusion MRI radiation necrosis tumor recurrence",
            "Galldiks PET radiation necrosis glioma",
        ],
        "requiredStudyTypes": ["guideline", "review"],
    },
    {
        "topic": "Rehabilitation after stroke: principles and evidence",
        "block": "Neurology",
        "priority": "high",
        "aliases": [
            "AVERT",
            "AHA ASA stroke rehabilitation",
            "early supported discharge stroke",
        ],
        "landmarkPmids": ["27145936", "25892679"],
        "guidelineQueries": [
            "AHA ASA adult stroke rehabilitation recovery guideline",
        ],
        "searchQueries": [
            "AVERT very early mobilisation within 24 hours stroke onset Lancet",
            "early supported discharge intensity multidisciplinary stroke rehabilitation Cochrane",
        ],
        "requiredStudyTypes": ["randomized controlled trial", "guideline"],
    },
    {
        "topic": "Retinal detachment: presentation and emergency management",
        "block": "Emergency Medicine",
        "priority": "medium",
        "aliases": [
            "rhegmatogenous retinal detachment",
            "pneumatic retinopexy",
            "RCOphth retinal detachment",
        ],
        "landmarkPmids": [],
        "guidelineQueries": [
            "Royal College of Ophthalmologists retinal detachment repair guideline",
        ],
        "searchQueries": [
            "pneumatic retinopexy versus scleral buckle vitrectomy rhegmatogenous retinal detachment",
            "emergency presentation retinal detachment management",
        ],
        "requiredStudyTypes": ["guideline", "review"],
    },
]


def lookup_pmid(title: str) -> tuple[str, int] | None:
    key = title.lower()
    for needle, value in TITLE_PMIDS.items():
        if needle in key:
            return value
    return None


def load_excel_topics() -> list[str]:
    wb = openpyxl.load_workbook(XLSX_SRC, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    topics = [str(r[0]).strip() for r in rows[1:] if r and r[0]]
    wb.close()
    return topics


def write_outputs(topics: list[str]) -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(XLSX_SRC, OUT_DIR / XLSX_SRC.name)
    index = OUT_DIR / "topics-index.txt"
    lines = [f"{i}. {t}" for i, t in enumerate(topics, 201)]
    index.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"Wrote {OUT_DIR}")


def enrich_corpus(topics: list[str]) -> None:
    corpus = json.loads(CORPUS_PATH.read_text(encoding="utf-8"))
    wanted = {t.lower() for t in topics}
    missing_topics = wanted - {t["topic"].lower() for t in corpus["topics"]}
    if missing_topics:
        raise SystemExit(f"Corpus missing topics: {sorted(missing_topics)}")
    enriched = 0
    for topic in corpus["topics"]:
        if topic["topic"].lower() not in wanted:
            continue
        for doc in topic["documents"]:
            if doc.get("pmid"):
                continue
            found = lookup_pmid(doc.get("title") or "")
            if not found:
                print(f"  NO PMID <- {(doc.get('title') or '')[:90]}")
                continue
            pmid, year = found
            doc["pmid"] = pmid
            doc["url"] = f"https://pubmed.ncbi.nlm.nih.gov/{pmid}/"
            if year and not doc.get("year"):
                doc["year"] = year
            enriched += 1
            print(f"  PMID {pmid} <- {doc['title'][:90]}")
    CORPUS_PATH.write_text(json.dumps(corpus, indent=2) + "\n", encoding="utf-8")
    print(f"Corpus PMID enrichments: {enriched}")


def add_alias(topic_row: dict, alias: str) -> None:
    aliases = topic_row.setdefault("aliases", [])
    if alias.lower() not in {a.lower() for a in aliases}:
        aliases.append(alias)


def add_pmid(topic_row: dict, pmid: str) -> None:
    pmids = topic_row.setdefault("landmarkPmids", [])
    if pmid not in pmids:
        pmids.append(pmid)


def update_flagship() -> None:
    config = json.loads(FLAGSHIP_PATH.read_text(encoding="utf-8"))
    by_name = {t["topic"].lower(): t for t in config["topics"]}

    pmr = by_name.get("polymyalgia rheumatica relapse prevention")
    if pmr:
        add_alias(pmr, "SAPHYR")
        add_alias(pmr, "SAPHROS")
        add_pmid(pmr, "37792612")
        add_pmid(pmr, "26359488")

    portal = by_name.get("portal hypertension and variceal bleeding")
    if portal:
        add_alias(portal, "PREDESCI")
        add_pmid(portal, "30910320")
        add_pmid(portal, "35120736")

    pa = by_name.get("primary aldosteronism diagnosis and treatment")
    if pa:
        add_alias(pa, "PATHWAY-2")
        add_pmid(pa, "26414968")
        add_pmid(pa, "26934393")

    mn = by_name.get("primary membranous nephropathy immunosuppression")
    if mn:
        add_pmid(mn, "31269364")
        add_pmid(mn, "34556256")

    icd = by_name.get("primary prevention icd sudden cardiac death")
    if icd:
        add_pmid(icd, "11907286")
        add_pmid(icd, "15659722")

    prl = by_name.get("prolactinoma dopamine agonist therapy")
    if prl:
        add_alias(prl, "Webster cabergoline")
        add_pmid(prl, "7915824")
        add_pmid(prl, "21296991")

    prone = by_name.get("prone positioning in severe ards")
    if prone:
        add_alias(prone, "prone positioning severe acute respiratory distress syndrome")
        add_pmid(prone, "23688302")

    pji = by_name.get("prosthetic joint infection diagnosis and retention strategy")
    if pji:
        add_alias(pji, "Zimmerli prosthetic joint")
        add_pmid(pji, "23223583")
        add_pmid(pji, "15483283")

    radial = by_name.get("radial versus femoral access for pci")
    if radial:
        add_pmid(radial, "21470671")
        add_pmid(radial, "25791214")

    pso = by_name.get("plaque psoriasis biologic therapy (il-17/il-23 inhibitors)")
    if pso:
        add_alias(pso, "Psoriasis: diagnosis, severity assessment and systemic treatment")
        add_pmid(pso, "30772098")
        add_pmid(pso, "28057361")
        add_pmid(pso, "31402114")

    sarc = by_name.get("sarcoidosis systemic immunosuppression")
    if sarc:
        add_alias(sarc, "Pulmonary sarcoidosis corticosteroid therapy")
        add_pmid(sarc, "16840744")

    lung = by_name.get("lung cancer screening low-dose ct")
    if lung:
        add_alias(lung, "pulmonary nodule evaluation and fleischner follow up")
        add_alias(lung, "Fleischner Society")
        add_pmid(lung, "21714641")
        add_pmid(lung, "31995683")
        add_pmid(lung, "28240562")

    existing = {t["topic"].lower() for t in config["topics"]}
    added = 0
    for row in NEW_FLAGSHIP:
        if row["topic"].lower() in existing:
            continue
        config["topics"].append(row)
        existing.add(row["topic"].lower())
        added += 1

    config["version"] = 18
    config["targets"]["flagshipCount"] = len(config["topics"])
    config["description"] = (
        f"Flagship-topic program - {len(config['topics'])} topics spanning evidence + learning surfaces, "
        "including curated literature cohorts Topics 51-75, Topics 76-100, Topics 126-150, "
        "Topics 151-175, Topics 176-200, and Topics 201-225."
    )
    FLAGSHIP_PATH.write_text(json.dumps(config, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"Flagship topics now {len(config['topics'])} (added {added})")


def main() -> None:
    if not XLSX_SRC.exists():
        raise SystemExit(f"Missing {XLSX_SRC}")
    topics = load_excel_topics()
    if len(topics) != 25:
        raise SystemExit(f"Expected 25 topics, found {len(topics)}")
    write_outputs(topics)
    enrich_corpus(topics)
    update_flagship()


if __name__ == "__main__":
    main()
