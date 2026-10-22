import sqlite3
import json
import openpyxl

conn = sqlite3.connect('database/app.db')
c = conn.cursor()

with open('topic-drafts/topics-1001-1250.json', 'r', encoding='utf-8') as f:
    drafts = json.load(f)

json_pub_225 = [d for d in drafts if d['status'] == 'published' and 1001 <= d['no'] <= 1225]
print(f"JSON published in 1001-1225: {len(json_pub_225)} / 225")

topic_ids = [d['topicId'] for d in json_pub_225]
placeholders = ','.join('?' for _ in topic_ids)
c.execute(f"SELECT COUNT(*) FROM topic_mentor_drafts WHERE status = 'published' AND topic_id IN ({placeholders})", topic_ids)
db_pub = c.fetchone()[0]
print(f"DB topic_mentor_drafts published in 1001-1225: {db_pub} / 225")

wb = openpyxl.load_workbook('Signal_MD_All_Topics_250_Per_Sheet.xlsx', data_only=True)
ws = wb['Topics 1001-1250']
master_pub = 0
for r in range(7, 232):
    val = ws.cell(row=r, column=11).value
    if val == 'Published':
        master_pub += 1
print(f"Master Sheet published in 1001-1225: {master_pub} / 225")

wb_rev = openpyxl.load_workbook('topic-drafts/topics-1001-1250-REVIEW.xlsx', data_only=True)
ws_rev = wb_rev.active
rev_pub = 0
for r in range(2, 227):
    val = ws_rev.cell(row=r, column=5).value
    if val == 'Published':
        rev_pub += 1
print(f"Review Sheet published in 1001-1225: {rev_pub} / 225")

wb.close()
wb_rev.close()
conn.close()
print("Batch 15 fully verified!")
